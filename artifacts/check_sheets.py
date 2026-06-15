import openpyxl

wb = openpyxl.load_workbook('artifacts/test_plan_subscribers.xlsx')
print(f"Sheets found: {wb.sheetnames}")
for name in wb.sheetnames:
    sheet = wb[name]
    print(f"\nSheet name: {name}")
    print(f"Dimensions: {sheet.dimensions}")
    # Print first 4 rows
    for r in range(1, 5):
        row_vals = [sheet.cell(r, c).value for c in range(1, 15)]
        print(f"Row {r}: {row_vals}")
